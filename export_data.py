import openpyxl, json

months = ['Jan/23','Feb/23','Mar/23','Apr/23','May/23','Jun/23','Jul/23','Aug/23','Sep/23','Oct/23','Nov/23','Dec/23',
          'Jan/24','Feb/24','Mar/24','Apr/24','May/24','Jun/24','Jul/24','Aug/24','Sep/24','Oct/24','Nov/24','Dec/24',
          'Jan/25','Feb/25','Mar/25','Apr/25','May/25','Jun/25','Jul/25','Aug/25','Sep/25','Oct/25','Nov/25','Dec/25']

state_fix = {
    'BAHIA': 'BA',
    'QUARTA PARADA Z/L': 'SP',
    'AC ': 'AC',
    'BA ': 'BA',
    'PB ': 'PB',
    'RS ': 'RS',
    'SP ': 'SP',
    'TO ': 'TO',
    'MG ': 'MG',
    '': 'N/A'
}

def extract_data(filepath, brand):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb['MRR_CLIENTES']
    rows = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if row[0] is None:
            continue
        categoria = str(row[1] or '').strip()
        if 'SAL' in categoria.upper():
            categoria = 'SALAO'
        estado = str(row[2] or '').strip().upper()
        estado = state_fix.get(estado, estado)
        status = str(row[4] or '').strip()
        distribuidor = str(row[5] or '').strip()
        
        monthly = {}
        for i, m in enumerate(months):
            val = row[6+i] if 6+i < len(row) else 0
            if val is None:
                val = 0
            if isinstance(val, str):
                val = val.replace('R$','').replace(' ','').strip()
                # Handle Brazilian format: 3.082,40 -> 3082.40
                if ',' in val:
                    val = val.replace('.','').replace(',','.')
                try:
                    val = float(val)
                except:
                    val = 0
            monthly[m] = round(float(val), 2)
        
        rows.append({
            'marca': brand,
            'categoria': categoria,
            'estado': estado,
            'status': status if status else 'ATIVO',
            'distribuidor': distribuidor,
            'monthly': monthly
        })
    return rows

visat = extract_data('MRR clientes VISAT Brasil 2023_2025..xlsx', 'VISAT')
prime = extract_data('MRR Distribuidores PRIME Brasil 2023_2025.xlsx', 'PRIME')

all_data = visat + prime

with open('data.json', 'w', encoding='utf-8') as f:
    json.dump(all_data, f, ensure_ascii=False, indent=None)

print(f"Exported {len(all_data)} records to data.json")
